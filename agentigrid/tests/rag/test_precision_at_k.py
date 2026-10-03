"""precision_at_k: parse C1's injected block, build the 15-row sheet, score it."""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location("precision_at_k", ROOT / "rag" / "tools" / "precision_at_k.py")
pk = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pk)

CTX = ("Use the following retrieved reference material only when relevant; do not invent facts beyond it.\n"
       "[ref 1 | score 0.71] Goal: lower cost\nStep applied: set_gen_status\n"
       "[ref 2 | score 0.55] opflow --help\n  -options\n"
       "[ref 3 | score 0.41] case list")


def test_parse_refs_keeps_multiline_chunks():
    refs = pk.parse_refs(CTX)
    assert [r["ref"] for r in refs] == [1, 2, 3]
    assert refs[0]["chunk"] == "Goal: lower cost\nStep applied: set_gen_status"
    assert refs[1]["chunk"].endswith("-options") and refs[2]["score"] == 0.41


def test_score_counts_missing_refs_as_not_relevant():
    rows = [{"goal_id": "a", "label": "1"}, {"goal_id": "a", "label": "0"}, {"goal_id": "a", "label": "1"},
            {"goal_id": "b", "label": "1"}, {"goal_id": "b", "label": ""}]
    res = pk.score(rows, 3)
    assert res["goals"]["a"]["precision_at_k"] == 0.667
    assert res["goals"]["b"]["precision_at_k"] == 0.333 and res["goals"]["b"]["precision_returned"] == 1.0
    assert res["unlabelled"] == 1 and res["macro_precision_at_k"] == 0.5


def test_sheet_round_trip(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    preview = {"goals": {"cost10": {"C1-basic": {"context": CTX}}}}
    spec = {"goals": [{"id": "cost10", "text": "Reduce cost"}]}
    items = pk.items_from_preview(preview, spec, "C1-basic")
    out = tmp_path / "p.xlsx"
    pk.write_sheet(items, out, {"preview_sha256": "x"})
    wb = load_workbook(out)
    ws = wb["Items"]
    assert [c.value for c in ws[1]] == pk.HEADERS and ws.max_row == 4
    assert "score" not in str([c.value for c in ws[2]]).lower()          # blind to similarity
    for i, v in zip((2, 3, 4), (1, 0, 1)):
        ws.cell(i, 5, v)
    wb.save(out)
    res = pk.score(pk.read_labels(out), 3)
    assert res["goals"]["cost10"]["relevant"] == 2
