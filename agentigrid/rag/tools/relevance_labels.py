#!/usr/bin/env python3
"""Relevance labels for the grader study (C2b thresholds, C2c calibration, H4).

Two steps, both offline (no LLM, no API):

  sample  Build a CSV of (query, chunk) pairs from the FROZEN, ingested corpus for
          a human to label 1 (the chunk would help an agent pursue this goal) or
          0 (it would not). Queries are the evaluated goal texts -- exactly what
          retrieval uses -- plus their CRAG keyword reformulation where it differs.
          Per query: the top --top chunks by cosine similarity (where grading
          decisions are made) and --random further chunks drawn at random (so
          clearly irrelevant pairs are represented too).

  fit     Read the labelled CSV, split it BY GOAL into a fitting set and a
          held-out set, then
            * C2b: fix the corrective thresholds on the reranker's unit-interval
              score (fitting set): tau_upper = lowest score at which precision of
              "relevant" >= --precision; tau_lower = highest score below which the
              share of irrelevant chunks >= --npv;
            * C2c: fit Platt scaling (2-parameter logistic regression) from the
              reranker's logit to P(relevant) on the fitting set;
            * H4 (held-out set): AUC of cosine, reranker and calibrated scores;
              expected calibration error (10 equal-width bins) of the reranker's
              unit score and of the calibrated probability.
          Writes a JSON report and the env lines for the C2b condition.

Usage (agentigrid project root, venv active, store ingested from the frozen corpus):
    OLLAMA_HOST=... python rag/tools/relevance_labels.py sample --out labels/pairs.csv
    # label the "label" column (1/0) in any spreadsheet tool, keep the CSV format
    python rag/tools/relevance_labels.py fit --labels labels/pairs.csv --heldout-goals relieve,n1cost10

Labelling by someone else (a domain expert) through a spreadsheet:
    python rag/tools/relevance_labels.py export-xlsx --pairs labels/pairs.csv --out labels/labeling_sheet.xlsx
    # send the .xlsx; the labeller fills "Label" (1/0/empty) and optionally "Comment"
    python rag/tools/relevance_labels.py import-xlsx --xlsx labels/labeling_sheet_filled.xlsx \
        --pairs labels/pairs.csv --out labels/pairs_labeled.csv
    python rag/tools/relevance_labels.py fit --labels labels/pairs_labeled.csv --heldout-goals relieve,n1cost10
  The sheet shows each (goal, chunk) pair once, in shuffled order, without cosine
  scores, ranks or query rewrites (blind to the retriever). import-xlsx copies each
  label to every sampled pair with that goal and chunk, and refuses a sheet made
  from a different pairs.csv.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import random
import sys
from datetime import datetime
from pathlib import Path

FIELDS = ["pair_id", "goal_id", "query_kind", "query", "doc_id", "source", "cosine", "rank",
          "label", "chunk"]


# ----------------------------------------------------------------------------- sample

def build_queries(spec: dict) -> list[dict]:
    from agentigrid.rag.corrective import CorrectiveRetriever
    reform = CorrectiveRetriever._reformulate
    out = []
    for g in spec.get("goals", []):
        out.append({"goal_id": g["id"], "kind": "goal", "text": g["text"]})
        r = reform(None, g["text"])
        if r != g["text"]:
            out.append({"goal_id": g["id"], "kind": "rewrite", "text": r})
    return out


def sample_pairs(queries, docs, query_fn, top: int, n_random: int, seed: int) -> list[dict]:
    """docs: [(id, text, meta)]; query_fn(text, k) -> [(doc_text, meta, cosine)]."""
    rng = random.Random(seed)
    by_text = {}
    for did, text, meta in docs:
        by_text.setdefault(text, (did, meta))
    rows = []
    for q in queries:
        hits = query_fn(q["text"], top)
        seen = set()
        for rank, (text, meta, cos) in enumerate(hits, 1):
            did, m = by_text.get(text, ("?", meta or {}))
            seen.add(did)
            rows.append(_row(q, did, (m or {}).get("source", ""), text, cos, rank))
        pool = [d for d in docs if d[0] not in seen]
        for did, text, meta in rng.sample(pool, min(n_random, len(pool))):
            rows.append(_row(q, did, (meta or {}).get("source", ""), text, None, "random"))
    for i, r in enumerate(rows, 1):
        r["pair_id"] = i
    return rows


def _row(q, did, source, text, cos, rank):
    return {"goal_id": q["goal_id"], "query_kind": q["kind"], "query": q["text"], "doc_id": did,
            "source": source, "cosine": "" if cos is None else round(float(cos), 4), "rank": rank,
            "label": "", "chunk": text}


def cmd_sample(args) -> int:
    import os
    from agentigrid.rag.corpus_hash import INGEST_MANIFEST_NAME
    from agentigrid.rag.embed import embed
    from agentigrid.rag.store import VectorStore

    spec = json.loads(Path(args.spec).read_text())
    host = os.environ.get("OLLAMA_HOST", "http://localhost:11434")
    store = VectorStore(args.store, args.collection, host, args.embed_model)
    docs = store.all_documents()
    if not docs:
        print(f"store {args.store} is empty -- ingest the frozen corpus first", file=sys.stderr)
        return 2
    queries = build_queries(spec)
    doc_emb = {}

    def cosine_all(text, k):
        # exact cosine over ALL chunks (not the ANN index), so ranks are reproducible
        qv = embed(text, host, args.embed_model)
        scored = []
        for did, d, meta in docs:
            if did not in doc_emb:
                doc_emb[did] = embed(d, host, args.embed_model)
            scored.append((d, meta, _cos(qv, doc_emb[did])))
        scored.sort(key=lambda x: -x[2])
        return scored[:k]

    rows = sample_pairs(queries, docs, cosine_all, args.top, args.random, args.seed)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    ing = Path(args.store) / INGEST_MANIFEST_NAME
    meta = {"created": datetime.now().isoformat(), "spec": args.spec, "store": args.store,
            "corpus_sha256": json.loads(ing.read_text()).get("corpus_sha256") if ing.exists() else None,
            "queries": len(queries), "pairs": len(rows), "top": args.top, "random": args.random,
            "seed": args.seed}
    out.with_suffix(".meta.json").write_text(json.dumps(meta, indent=2) + "\n")
    print(f"Wrote {len(rows)} pairs for {len(queries)} queries -> {out}\n"
          f"Label the 'label' column 1 (would help an agent pursue this goal) or 0 (would not).")
    return 0


def _cos(a, b):
    na = math.sqrt(sum(x * x for x in a)) or 1.0
    nb = math.sqrt(sum(x * x for x in b)) or 1.0
    return sum(x * y for x, y in zip(a, b)) / (na * nb)


# ----------------------------------------------------------------------------- fit

def auc(scores, labels) -> float | None:
    """Mann-Whitney AUC; ties count 1/2. None if a class is missing."""
    pos = [s for s, y in zip(scores, labels) if y == 1]
    neg = [s for s, y in zip(scores, labels) if y == 0]
    if not pos or not neg:
        return None
    wins = 0.0
    for p in pos:
        for n in neg:
            wins += 1.0 if p > n else 0.5 if p == n else 0.0
    return wins / (len(pos) * len(neg))


def ece(probs, labels, bins: int = 10) -> float | None:
    if not probs:
        return None
    total, err = len(probs), 0.0
    for b in range(bins):
        lo, hi = b / bins, (b + 1) / bins
        idx = [i for i, p in enumerate(probs) if (lo <= p < hi) or (b == bins - 1 and p == 1.0)]
        if idx:
            conf = sum(probs[i] for i in idx) / len(idx)
            acc = sum(labels[i] for i in idx) / len(idx)
            err += len(idx) / total * abs(acc - conf)
    return err


def platt_fit(x, y, iters: int = 200):
    """Platt scaling: P(y=1|x) = 1 / (1 + exp(-(a*x + b))), with Platt's target
    smoothing, fitted by Newton's method. Returns (a, b)."""
    n1 = sum(y)
    n0 = len(y) - n1
    t1, t0 = (n1 + 1) / (n1 + 2), 1 / (n0 + 2)
    t = [t1 if yi == 1 else t0 for yi in y]
    a, b = 0.0, math.log((n0 + 1) / (n1 + 1))
    for _ in range(iters):
        g_a = g_b = h_aa = h_ab = h_bb = 0.0
        for xi, ti in zip(x, t):
            p = 1.0 / (1.0 + math.exp(-(a * xi + b)))
            d = p - ti
            w = p * (1 - p)
            g_a += d * xi; g_b += d
            h_aa += w * xi * xi; h_ab += w * xi; h_bb += w
        det = h_aa * h_bb - h_ab * h_ab
        if abs(det) < 1e-12:
            break
        da = (h_bb * g_a - h_ab * g_b) / det
        db = (h_aa * g_b - h_ab * g_a) / det
        a, b = a - da, b - db
        if abs(da) < 1e-9 and abs(db) < 1e-9:
            break
    return a, b


def thresholds(scores, labels, precision: float, npv: float):
    """tau_upper: lowest threshold with precision(score >= tau) >= `precision`;
    tau_lower: highest threshold with share of 0-labels among score < tau >= `npv`."""
    cand = sorted(set(scores))
    tau_u = None
    for c in cand:
        sel = [y for s, y in zip(scores, labels) if s >= c]
        if sel and sum(sel) / len(sel) >= precision:
            tau_u = c
            break
    tau_l = None
    for c in reversed(cand):
        sel = [y for s, y in zip(scores, labels) if s < c]
        if sel and (len(sel) - sum(sel)) / len(sel) >= npv:
            tau_l = c
            break
    if tau_u is not None and tau_l is not None and tau_l > tau_u:
        tau_l = tau_u
    return tau_l, tau_u


def _logit(p, eps=1e-6):
    p = min(max(p, eps), 1 - eps)
    return math.log(p / (1 - p))


def fit_report(rows, heldout: set, reranker_fn, precision=0.8, npv=0.8) -> dict:
    lab = [r for r in rows if str(r.get("label", "")).strip() in ("0", "1")]
    if not lab:
        raise SystemExit("no labelled rows (label column must be 0 or 1)")
    for r in lab:
        r["y"] = int(str(r["label"]).strip())
    missing = heldout - {r["goal_id"] for r in lab}
    if missing:
        raise SystemExit(f"held-out goals without labels: {sorted(missing)}")
    scores = reranker_fn([(r["query"], r["chunk"]) for r in lab])   # unit interval
    for r, s in zip(lab, scores):
        r["rerank"] = float(s)
    fit = [r for r in lab if r["goal_id"] not in heldout]
    test = [r for r in lab if r["goal_id"] in heldout]
    if not fit or not test:
        raise SystemExit("need labelled rows in both the fitting and the held-out goals")
    tau_l, tau_u = thresholds([r["rerank"] for r in fit], [r["y"] for r in fit], precision, npv)
    a, b = platt_fit([_logit(r["rerank"]) for r in fit], [r["y"] for r in fit])

    def calib(p):
        return 1.0 / (1.0 + math.exp(-(a * _logit(p) + b)))

    ty = [r["y"] for r in test]
    cos_ok = [r for r in test if str(r.get("cosine", "")) not in ("", "None")]
    report = {
        "created": datetime.now().isoformat(),
        "n_labelled": len(lab), "n_fit": len(fit), "n_heldout": len(test),
        "heldout_goals": sorted(heldout), "positives_fit": sum(r["y"] for r in fit),
        "positives_heldout": sum(ty),
        "rule": {"precision": precision, "npv": npv},
        "C2b": {"tau_lower": tau_l, "tau_upper": tau_u,
                "env": {"AGENTIGRID_CRAG_TAU_LOWER": tau_l, "AGENTIGRID_CRAG_TAU_UPPER": tau_u}},
        "C2c": {"platt_a": a, "platt_b": b},
        "heldout": {
            "auc_cosine": auc([float(r["cosine"]) for r in cos_ok], [r["y"] for r in cos_ok]),
            "auc_cosine_note": "top-k pairs only (random pairs carry no cosine rank)",
            "auc_reranker": auc([r["rerank"] for r in test], ty),
            "auc_calibrated": auc([calib(r["rerank"]) for r in test], ty),
            "ece_reranker_unit": ece([r["rerank"] for r in test], ty),
            "ece_calibrated": ece([calib(r["rerank"]) for r in test], ty),
        },
    }
    return report


def cmd_fit(args) -> int:
    with open(args.labels, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    heldout = {g.strip() for g in args.heldout_goals.split(",") if g.strip()}
    from agentigrid.rag.grader_reranker import RerankerGrader
    grader = RerankerGrader(model_name=args.model) if args.model else RerankerGrader()

    def rerank(pairs):
        hits = [(chunk, {}, 0.0) for _q, chunk in pairs]
        out = []
        for (q, _c), h in zip(pairs, hits):        # one query per pair
            out.extend(grader(q, [h]))
        return out

    report = fit_report(rows, heldout, rerank, args.precision, args.npv)
    report["reranker"] = grader.describe()
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    Path(args.out).write_text(json.dumps(report, indent=2) + "\n")
    c = report["C2b"]
    print(json.dumps(report["heldout"], indent=2))
    print(f"\nC2b thresholds (fitting goals): tau_lower={c['tau_lower']}, tau_upper={c['tau_upper']}")
    print(f"C2c Platt: a={report['C2c']['platt_a']:.4f}, b={report['C2c']['platt_b']:.4f}")
    print(f"Report -> {args.out}")
    return 0


# ----------------------------------------------------------------------------- labelling sheet

SHEET_PAIRS, SHEET_GOALS, SHEET_INFO, SHEET_META = "Pairs", "Goals", "Instructions", "_meta"
PAIR_HEADERS = ["Item", "Goal id", "Goal", "Chunk", "Label (1/0)", "Comment", "Key"]
LABELLER_FIELDS = ("Name", "Role / affiliation", "Date completed")

INSTRUCTIONS = [
    ("h", "Relevance labelling: AgentiGrid retrieval study"),
    ("p", "An AI agent (a large language model) operates the ExaGO power-grid optimization toolkit "
          "to reach a goal stated in plain language, for example reducing generation cost or relieving "
          "line overloads. Before each step it can be shown short passages (\"chunks\") retrieved from a "
          "reference collection: ExaGO tool help, AgentiGrid command formats, and solver-verified worked "
          "examples from other grid cases."),
    ("p", "Your labels say which chunks are relevant to which goal. They are the ground truth against "
          "which the retrieval graders are evaluated, and they are reported in a research paper."),
    ("s", "Your task"),
    ("p", "For each row of the Pairs sheet, read the Goal and the Chunk, then enter 1 or 0 in the yellow "
          "Label column. The Goals sheet describes each goal, the grid case and the application it uses."),
    ("p", "The question to ask: if the agent read only this chunk, would its next proposal for this goal "
          "be more likely to be valid and on task?"),
    ("s", "Label 1 when the chunk"),
    ("b", "shows a worked step of the same kind of task, for example a solver-verified cost reduction by "
          "unit commitment for a cost goal;"),
    ("b", "explains a command or option the goal needs, for example generator dispatch or commitment for "
          "cost; transformer tap, shunt or generator voltage setpoint for voltage; redispatch for overload "
          "relief;"),
    ("b", "gives a tool fact the goal depends on, for example SCOPFLOW contingency options for the N-1 goal."),
    ("s", "Label 0 when the chunk"),
    ("b", "is about a different kind of task, for example a voltage example for a cost goal;"),
    ("b", "covers an application the goal does not use (see the Goals sheet);"),
    ("b", "is a generic listing, for example a list of case files;"),
    ("b", "shows an action the goal forbids, for example changing voltage limits for a goal that says not to."),
    ("s", "Notes"),
    ("b", "Judge relevance to the goal, not the quality of the chunk. Worked examples come from other "
          "networks (case9mod, case118, ACTIVSg500) while the goals concern ACTIVSg200 and case39; an "
          "example from another network can still be relevant, because the method transfers even though "
          "bus numbers do not."),
    ("b", "If you cannot decide, leave Label empty and say why in Comment. Please keep empty labels to a "
          "small share of the rows."),
    ("b", "Please label from your own judgment, without AI tools: the labels must be an independent human "
          "reference."),
    ("b", "Edit only the yellow cells. Sorting and filtering the Pairs sheet is fine; do not delete rows, "
          "edit other columns, or rename sheets. Save as .xlsx and send the file back."),
]
EXAMPLE = ("cost10",
           "Goal: Lower the total generation cost by at least 5% while serving the same load, without "
           "changing generator cost curves, and without relaxing bus voltage limits or line ratings\n"
           "Step applied: set_gen_status x4 on buses 435, 436, 125, 126 (status=1)\n"
           "Result: OPFLOW converged ...",
           1, "Same kind of task (cost reduction by unit commitment)")


def _sha256(path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _goal_info(spec: dict) -> list[dict]:
    """One row per evaluated goal: text, application(s), case(s), success criterion."""
    where = {}
    for c in spec.get("cases", []):
        for gid in c.get("goals", []):
            where.setdefault(gid, []).append((c.get("app", ""), c.get("name", "")))
    out = []
    for g in spec.get("goals", []):
        if g.get("target_pct") is not None and g.get("success") == "n1_secure":
            crit = f"dispatch secure against all listed contingencies AND cost reduced by at least {g['target_pct']}%"
        elif g.get("target_pct") is not None:
            crit = f"cost reduced by at least {g['target_pct']}% versus the starting point"
        elif g.get("success") == "no_violations":
            crit = "power flow with no voltage or line-loading violations"
        else:
            crit = "no pass/fail criterion (the agent searches for a maximum)"
        places = where.get(g["id"], [])
        out.append({"id": g["id"], "text": g["text"],
                    "apps": ", ".join(sorted({a.upper() for a, _ in places})),
                    "cases": "; ".join(n for _, n in places), "criterion": crit})
    return out


def group_items(rows: list[dict]) -> list[dict]:
    """Unique (goal, chunk) items; each keeps the pair_ids it stands for."""
    items = {}
    for r in rows:
        key = (r["goal_id"], hashlib.sha256(r["chunk"].encode("utf-8")).hexdigest())
        it = items.setdefault(key, {"goal_id": r["goal_id"], "chunk": r["chunk"], "pair_ids": []})
        it["pair_ids"].append(str(r["pair_id"]))
    return list(items.values())


def _row_height(text: str, chars_per_line: int) -> float:
    lines = sum(max(1, math.ceil(len(line) / chars_per_line)) for line in text.split("\n"))
    return min(409.0, 15.0 * lines + 4)


def export_xlsx(pairs_path, spec_path, out_path, seed: int = 11) -> dict:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    with open(pairs_path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    spec = json.loads(Path(spec_path).read_text())
    goals = _goal_info(spec)
    goal_text = {g["id"]: g["text"] for g in goals}
    items = group_items(rows)
    random.Random(seed).shuffle(items)

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head_fill = PatternFill("solid", start_color="D9E1F2")
    input_fill = PatternFill("solid", start_color="FFFF00")
    wrap_top = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    # --- Instructions
    ws = wb.active
    ws.title = SHEET_INFO
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    r = 1
    for kind, text in INSTRUCTIONS:
        if kind == "h":
            ws.cell(r, 1, text).font = Font(name="Arial", size=14, bold=True)
            r += 2
            continue
        if kind == "s":
            r += 1
            ws.cell(r, 1, text).font = Font(name="Arial", size=11, bold=True)
        else:
            c = ws.cell(r, 2, ("• " if kind == "b" else "") + text)
            c.font, c.alignment = font, wrap_top
            ws.row_dimensions[r].height = _row_height(c.value, 100)
        r += 1
    r += 1
    ws.cell(r, 1, "Legend").font = Font(name="Arial", size=11, bold=True)
    r += 1
    leg = ws.cell(r, 1, "")
    leg.fill = input_fill
    ws.cell(r, 2, "Yellow cells are the only cells to fill in (Label and Comment on the Pairs sheet, "
                  "and your details below).").font = font
    r += 2
    ws.cell(r, 1, "Example (not part of the Pairs sheet)").font = Font(name="Arial", size=11, bold=True)
    r += 1
    for label, val in zip(("Goal id", "Chunk (shortened)", "Label (1/0)", "Comment"), EXAMPLE):
        lc = ws.cell(r, 1, label)
        lc.font, lc.alignment = bold, Alignment(vertical="top")
        c = ws.cell(r, 2, val)
        c.font, c.alignment = font, wrap_top
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        ws.row_dimensions[r].height = _row_height(str(val), 100)
        r += 1
    r += 1
    ws.cell(r, 1, "Progress").font = Font(name="Arial", size=11, bold=True)
    r += 1
    n = len(items)
    lab_col = f"'{SHEET_PAIRS}'!$E$2:$E${n + 1}"
    progress = [("Items", f"=COUNTA('{SHEET_PAIRS}'!$A$2:$A${n + 1})"),
                ("Labelled 1", f"=COUNTIF({lab_col},1)"),
                ("Labelled 0", f"=COUNTIF({lab_col},0)")]
    first = r
    for label, formula in progress:
        ws.cell(r, 1, label).font = bold
        c = ws.cell(r, 2, formula)
        c.font, c.alignment = font, Alignment(horizontal="left")
        r += 1
    ws.cell(r, 1, "Still empty").font = bold
    c = ws.cell(r, 2, f"=B{first}-B{first + 1}-B{first + 2}")
    c.font, c.alignment = font, Alignment(horizontal="left")
    r += 2
    ws.cell(r, 1, "Your details").font = Font(name="Arial", size=11, bold=True)
    r += 1
    labeller_cells = {}
    for field in LABELLER_FIELDS:
        ws.cell(r, 1, field).font = bold
        c = ws.cell(r, 2, None)
        c.fill, c.border, c.font = input_fill, box, font
        labeller_cells[field] = f"B{r}"
        r += 1

    # --- Goals
    wg = wb.create_sheet(SHEET_GOALS)
    heads = ["Goal id", "Goal (exact text the agent receives)", "Application", "Grid case(s)", "Success criterion"]
    widths = [12, 70, 14, 34, 40]
    for j, (h, w) in enumerate(zip(heads, widths), 1):
        c = wg.cell(1, j, h)
        c.font, c.fill, c.alignment, c.border = bold, head_fill, wrap_top, box
        wg.column_dimensions[c.column_letter].width = w
    for i, g in enumerate(goals, 2):
        for j, v in enumerate((g["id"], g["text"], g["apps"], g["cases"], g["criterion"]), 1):
            c = wg.cell(i, j, v)
            c.font, c.alignment, c.border = font, wrap_top, box
        wg.row_dimensions[i].height = _row_height(g["text"], 66)
    wg.freeze_panes = "A2"
    wg.cell(len(goals) + 3, 1, "Application: OPFLOW = AC optimal power flow; SCOPFLOW = security-constrained "
                               "OPF (N-1 contingencies); PFLOW = power flow (fixed set-points, limits "
                               "reported but not enforced).").font = font

    # --- Pairs
    wp = wb.create_sheet(SHEET_PAIRS)
    widths = [7, 10, 45, 100, 12, 30, 12]
    for j, (h, w) in enumerate(zip(PAIR_HEADERS, widths), 1):
        c = wp.cell(1, j, h)
        c.font, c.fill, c.alignment, c.border = bold, head_fill, wrap_top, box
        wp.column_dimensions[c.column_letter].width = w
    wp.cell(1, 5).fill = input_fill
    wp.cell(1, 6).fill = input_fill
    wp.cell(1, 5).comment = Comment("Enter 1 (relevant) or 0 (not relevant); leave empty if you cannot decide.",
                                    "AgentiGrid")
    for i, it in enumerate(items, 2):
        vals = (i - 1, it["goal_id"], goal_text.get(it["goal_id"], ""), it["chunk"], None, None,
                ";".join(it["pair_ids"]))
        for j, v in enumerate(vals, 1):
            c = wp.cell(i, j, v)
            c.font, c.alignment, c.border = font, wrap_top, box
        wp.cell(i, 5).fill = input_fill
        wp.cell(i, 5).alignment = Alignment(horizontal="center", vertical="top")
        wp.cell(i, 6).fill = input_fill
        wp.row_dimensions[i].height = max(_row_height(it["chunk"], 95),
                                          _row_height(goal_text.get(it["goal_id"], ""), 42))
    dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Label", error="Enter 1, 0, or leave the cell empty.")
    wp.add_data_validation(dv)
    dv.add(f"E2:E{n + 1}")
    wp.column_dimensions["G"].hidden = True
    wp.freeze_panes = "C2"
    wp.auto_filter.ref = f"A1:F{n + 1}"

    # --- meta (hidden)
    wm = wb.create_sheet(SHEET_META)
    meta = {"pairs_file": str(pairs_path), "pairs_sha256": _sha256(pairs_path), "pairs": len(rows),
            "items": n, "seed": seed, "spec": str(spec_path), "created": datetime.now().isoformat(),
            "labeller_cells": json.dumps(labeller_cells)}
    for i, (k, v) in enumerate(meta.items(), 1):
        wm.cell(i, 1, k)
        wm.cell(i, 2, v)
    wm.sheet_state = "hidden"
    for sheet, landscape in ((ws, False), (wg, True), (wp, True)):   # print: fit columns to page width
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    wp.print_title_rows = "1:1"
    wb.active = 0
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {"pairs": len(rows), "items": n, "goals": len(goals), "out": str(out_path)}


def _parse_label(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return ""
    try:
        f = float(str(v).strip())
    except ValueError:
        raise ValueError(f"label {v!r} is not 1, 0 or empty")
    if f not in (0.0, 1.0):
        raise ValueError(f"label {v!r} is not 1, 0 or empty")
    return str(int(f))


def import_xlsx(xlsx_path, pairs_path, out_path, force: bool = False) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, data_only=True)
    for name in (SHEET_PAIRS, SHEET_META, SHEET_INFO):
        if name not in wb.sheetnames:
            raise SystemExit(f"{xlsx_path}: sheet {name!r} missing -- not a sheet made by export-xlsx")
    meta = {row[0]: row[1] for row in wb[SHEET_META].iter_rows(values_only=True) if row[0]}
    digest = _sha256(pairs_path)
    if meta.get("pairs_sha256") != digest and not force:
        raise SystemExit(f"{pairs_path} is not the pairs file this sheet was made from "
                         f"(sha256 differs); pass the original, or --force")
    ws = wb[SHEET_PAIRS]
    header = [c.value for c in ws[1]]
    try:
        col = {h: header.index(h) for h in ("Label (1/0)", "Comment", "Key")}
    except ValueError as exc:
        raise SystemExit(f"Pairs sheet header changed ({exc}); expected {PAIR_HEADERS}")
    labels, comments, bad = {}, {}, []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        key = row[col["Key"]]
        if not key:
            continue
        try:
            lab = _parse_label(row[col["Label (1/0)"]])
        except ValueError as exc:
            bad.append(f"row {i}: {exc}")
            continue
        com = row[col["Comment"]] or ""
        for pid in str(key).split(";"):
            labels[pid], comments[pid] = lab, str(com).strip()
    if bad:
        raise SystemExit("invalid labels:\n  " + "\n  ".join(bad))
    with open(pairs_path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    missing = [r["pair_id"] for r in rows if str(r["pair_id"]) not in labels]
    if missing and not force:
        raise SystemExit(f"{len(missing)} pair(s) not found in the sheet (rows deleted?), e.g. {missing[:5]}")
    per_goal = {}
    for r in rows:
        pid = str(r["pair_id"])
        r["label"] = labels.get(pid, "")
        r["comment"] = comments.get(pid, "")
        g = per_goal.setdefault(r["goal_id"], {"1": 0, "0": 0, "": 0})
        g[r["label"]] += 1
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS + ["comment"])
        w.writeheader()
        w.writerows(rows)
    info = wb[SHEET_INFO]
    cells = json.loads(meta.get("labeller_cells") or "{}")
    labeller = {k: (info[v].value if v else None) for k, v in cells.items()}
    side = {"source_sheet": str(xlsx_path), "sheet_sha256": _sha256(xlsx_path), "pairs_file": str(pairs_path),
            "pairs_sha256": digest, "imported": datetime.now().isoformat(), "labeller": labeller,
            "per_goal": per_goal}
    Path(out_path).with_suffix(".labeller.json").write_text(json.dumps(side, indent=2, default=str) + "\n")
    return side


def cmd_export_xlsx(args) -> int:
    res = export_xlsx(args.pairs, args.spec, args.out, args.seed)
    print(f"Wrote {res['items']} unique (goal, chunk) item(s) covering {res['pairs']} pair(s) -> {res['out']}")
    return 0


def cmd_import_xlsx(args) -> int:
    side = import_xlsx(args.xlsx, args.pairs, args.out, args.force)
    for g, c in sorted(side["per_goal"].items()):
        print(f"  {g:<10} 1={c['1']:<4} 0={c['0']:<4} empty={c['']}")
    print(f"Labeller: {side['labeller']}")
    print(f"Wrote {args.out} (and {Path(args.out).with_suffix('.labeller.json')})")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample")
    s.add_argument("--spec", default="grader_ablation_spec.json")
    s.add_argument("--store", default="rag/store")
    s.add_argument("--collection", default="agentigrid_kb")
    s.add_argument("--embed-model", default="nomic-embed-text")
    s.add_argument("--top", type=int, default=25)
    s.add_argument("--random", type=int, default=10)
    s.add_argument("--seed", type=int, default=7)
    s.add_argument("--out", default="labels/pairs.csv")
    f = sub.add_parser("fit")
    f.add_argument("--labels", required=True)
    f.add_argument("--heldout-goals", required=True, help="comma-separated goal ids held out for AUC/ECE")
    f.add_argument("--precision", type=float, default=0.8)
    f.add_argument("--npv", type=float, default=0.8)
    f.add_argument("--model", default=None, help="reranker model (default: grader default / env)")
    f.add_argument("--out", default="labels/fit_report.json")
    e = sub.add_parser("export-xlsx", help="labelling spreadsheet (with instructions) for a human labeller")
    e.add_argument("--pairs", default="labels/pairs.csv")
    e.add_argument("--spec", default="grader_ablation_spec.json")
    e.add_argument("--out", default="labels/labeling_sheet.xlsx")
    e.add_argument("--seed", type=int, default=11, help="row shuffle seed")
    i = sub.add_parser("import-xlsx", help="read a filled labelling sheet back into a labelled CSV")
    i.add_argument("--xlsx", required=True)
    i.add_argument("--pairs", default="labels/pairs.csv")
    i.add_argument("--out", default="labels/pairs_labeled.csv")
    i.add_argument("--force", action="store_true", help="accept a pairs-file mismatch or missing rows")
    args = ap.parse_args(argv)
    return {"sample": cmd_sample, "fit": cmd_fit, "export-xlsx": cmd_export_xlsx,
            "import-xlsx": cmd_import_xlsx}[args.cmd](args)


if __name__ == "__main__":
    raise SystemExit(main())
